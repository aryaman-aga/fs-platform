from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pytest
from helpers_report import XSS, make_applied, make_finding, make_plan, make_result

from fscore.report.common import finding_status, register_filename
from fscore.report.register import (
    CHANGE_HEADERS,
    EVIDENCE_HEADERS,
    EXCEPTION_HEADERS,
    KEY_HEADERS,
    MOVE_HEADERS,
    PDF_HEADERS,
    XREF_HEADERS,
    write_register,
)

SHEETS = ["Exceptions", "Evidence", "Formula changes", "Printed figures moved", "PDF reconciliation", "Cross-reference map", "Key figures", "Run info"]


def _headers(ws: openpyxl.worksheet.worksheet.Worksheet, n: int) -> list[str]:
    return [ws.cell(1, c).value for c in range(1, n + 1)]


def test_register_has_all_sheets_and_headers(out_dir: Path) -> None:
    path = write_register(make_result(), out_dir / "reg.xlsx", plan=make_plan(), applied=make_applied())
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == SHEETS
    for name, headers in [
        ("Exceptions", EXCEPTION_HEADERS), ("Evidence", EVIDENCE_HEADERS), ("Formula changes", CHANGE_HEADERS),
        ("Printed figures moved", MOVE_HEADERS), ("PDF reconciliation", PDF_HEADERS), ("Cross-reference map", XREF_HEADERS),
        ("Key figures", KEY_HEADERS),
    ]:
        ws = wb[name]
        assert _headers(ws, len(headers)) == headers, name
        assert ws.freeze_panes == "A2", name
        assert ws.auto_filter.ref, name


def test_exceptions_rows_sorted_and_statuses(out_dir: Path) -> None:
    result = make_result()
    path = write_register(result, out_dir / "reg.xlsx", plan=make_plan(), applied=make_applied())
    ws = openpyxl.load_workbook(path)["Exceptions"]
    rows = [[ws.cell(r, c).value for c in range(1, len(EXCEPTION_HEADERS) + 1)] for r in range(2, ws.max_row + 1)]
    assert len(rows) == len(result.findings)
    priorities = [r[1].split(" / ")[0] for r in rows]
    assert priorities == sorted(priorities, key=["High", "Medium", "Low", "Info"].index)
    by_rule = {r[2].split()[0]: r for r in rows}
    assert by_rule["CAST-01"][3] == "Balance Sheet"  # statement label from the statements map
    assert by_rule["CAST-01"][9] == pytest.approx(0.01)
    assert by_rule["CAST-01"][11] == "Fixed"
    assert by_rule["TXT-01"][11] == "Proposed – needs approval"
    assert by_rule["RPT-09"][11] == "Open – manual"
    assert by_rule["TXT-07"][3] == "Note 4-10"  # unknown sheets fall back to the sheet name


def test_changes_moves_and_pdf_sheets(out_dir: Path) -> None:
    path = write_register(make_result(), out_dir / "reg.xlsx", plan=make_plan(), applied=make_applied())
    wb = openpyxl.load_workbook(path)
    ch = wb["Formula changes"]
    assert ch.cell(2, 6).value == "=SUM(C20:C21)"
    assert ch.cell(2, 7).value == "=SUMPRODUCT(ROUND(C20:C21,2))"
    assert ch.cell(2, 13).value == "Applied"
    assert ch.cell(4, 13).value == "Skipped"
    mv = wb["Printed figures moved"]
    assert mv.max_row == 4
    assert mv.cell(3, 7).value == "Yes"  # PY flagged
    pdf = wb["PDF reconciliation"]
    texts = [str(pdf.cell(r, 6).value or "") for r in range(2, pdf.max_row + 1)]
    assert any("older version" in t for t in texts)
    assert any("re-print" in t.lower() for t in texts)
    info = wb["Run info"]
    items = {info.cell(r, 1).value: info.cell(r, 2).value for r in range(2, info.max_row + 1)}
    assert items["File SHA-256"] == "cd" * 32
    assert items["Fix plan SHA-256"] == make_plan().plan_sha256
    assert items["Post-write verification"] == "Passed"


def test_register_without_plan_or_pdf(out_dir: Path) -> None:
    path = write_register(make_result(findings=[], pdf=False), out_dir / "empty.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb["Exceptions"].cell(2, 7).value == "No findings."
    assert "No printed PDF" in wb["PDF reconciliation"].cell(2, 6).value


def test_text_is_never_written_as_formula(out_dir: Path) -> None:
    f = make_finding("F-1", explanation="=HYPERLINK(\"http://example.invalid\",\"x\") " + XSS)
    path = write_register(make_result(findings=[f], pdf=False), out_dir / "r.xlsx")
    ws = openpyxl.load_workbook(path)["Exceptions"]
    cell = ws.cell(2, 7)
    assert cell.data_type == "s"
    assert cell.value.startswith("=HYPERLINK")


def test_register_filename_is_safe() -> None:
    assert register_filename("Corvus Data: <Services>/FY26", date(2026, 10, 21)) == "Corvus Data Services FY26 FS-Check 2026-10-21.xlsx"
    assert register_filename("", date(2026, 1, 2)) == "Client FS-Check 2026-01-02.xlsx"


def test_finding_status_without_plan() -> None:
    f = make_finding("F-9")
    assert finding_status(f, None, None) == "Open"
