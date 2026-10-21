"""Module P output on the synthetic books: a clean, conventional Ind AS workbook."""

from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest
from helpers_synthetic import cached_workbook, printed_alias

from fscore.engine.calc import Recalculator
from fscore.engine.evaluator import values_equal
from fscore.io.xlsx_reader import Workbook
from fscore.prepare.statements import SHEET_ORDER
from fscore.synthetic.oracle import print_sheets, scan_casting, simulate_totals_fix, tie_gaps

PERIODS = ("CY", "PY")


@pytest.fixture(scope="module")
def clean():
    return cached_workbook(seed=7)


@pytest.fixture(scope="module")
def wb(clean):
    w = Workbook.open(clean.path)
    yield w
    w.close()


def test_sheet_order_print_group_and_print_areas(wb: Workbook) -> None:
    assert wb.sheet_names == SHEET_ORDER
    group = print_sheets(wb)
    assert group == SHEET_ORDER[:12]
    for name in group:
        s = wb.sheet(name)
        assert s.print_area, f"{name} has a print area"
    for name in ("NS_TB_FT27", "FAR", "Ageing"):
        assert not wb.sheet(name).tab_selected


def test_balance_sheet_tallies_on_print(clean, wb: Workbook) -> None:
    for k in PERIODS:
        assert printed_alias(wb, clean.index, "bs.total_assets", k) == printed_alias(wb, clean.index, "bs.total_eq_liab", k)
    assert tie_gaps(clean.path, clean.index) == {"bs": {"CY": 0, "PY": 0}, "cf": {"CY": 0, "PY": 0}}


def test_every_printed_total_casts(clean) -> None:
    assert scan_casting(clean.path) == []
    assert clean.expected == []
    assert clean.gaps == {"bs": {"CY": 0, "PY": 0}, "cf": {"CY": 0, "PY": 0}}


def test_key_ties(clean, wb: Workbook) -> None:
    ix = clean.index
    p = lambda a, k: printed_alias(wb, ix, a, k)  # noqa: E731
    for k in PERIODS:
        assert p("cf.closing_cash", k) == p("bs.cash", k) == p("note.cash.total", k) == p("cf.components_total", k)
        assert p("note.other_equity.total", k) == p("bs.other_equity", k)
        assert p("note.ppe.net_cy" if k == "CY" else "note.ppe.net_py", "total") == p("bs.ppe", k)
        assert p("note.tr.ageing_total_" + k.lower(), "total") == p("note.tr.gross", k)
        assert p("note.tp.ageing_total_" + k.lower(), "total") == p("note.tp.total", k)
        assert p("note.dbo.close", k) == p("note.dbo.split_total", k)
        assert p("note.dbo.expense", k) == p("note.employee.gratuity", k)
        assert p("note.revenue.disaggregation_total", k) == p("note.revenue.total", k)
        assert p("note.tax.rec_total", k) == p("note.tax.total", k)
        assert p("note.ctl.close", k) == p("note.ctl.bs", k)
        assert p("note.lease.close", k) == p("note.lease.total", k)
        assert p("note.auditor.total", k) == p("note.other_expenses.auditor", k)
        assert p("note.msme.principal", k) == p("note.tp.msme", k)
    # SOCE closing = BS equity
    s, r, cols = ix["soce.close_cy"]
    assert wb.render_cell(wb.cell(s, r, cols["total"])).value == p("bs.total_equity", "CY")
    s, r, cols = ix["soce.close_py"]
    assert wb.render_cell(wb.cell(s, r, cols["total"])).value == p("bs.total_equity", "PY")
    # PAT = PBT − tax; TCI = PAT + OCI
    for k in PERIODS:
        assert p("pl.pat", k) == p("pl.pbt", k) - p("pl.tax_total", k)
        assert p("pl.tci", k) == p("pl.pat", k) + p("pl.oci", k)


def test_balancing_lines_exist_with_house_captions(clean, wb: Workbook) -> None:
    s, r, _ = clean.index["note.other_expenses.misc"]
    assert wb.cell(s, r, 1).value == "Miscellaneous expenses"
    s, r, cols = clean.index["pl.other_expenses"]
    assert "Note 20-27" in (wb.cell(s, r, cols["CY"]).formula or "")
    s, r, _ = clean.index["cf.ocl"]
    assert wb.cell(s, r, 1).value == "(Decrease)/increase in other current liabilities"


def test_totals_are_formulas_everywhere(wb: Workbook) -> None:
    """No hard-typed totals in the clean fixture (TYPED-01), PY / typed columns included."""
    for name in print_sheets(wb):
        s = wb.sheet(name)
        for c in s.iter_cells():
            if c.col != 1 or not isinstance(c.value, str):
                continue
            label = c.value.strip().lower()
            if label.startswith(("total", "balance at the end", "balance as at")) or label in ("total assets", "total equity and liabilities"):
                for col in range(2, s.max_col + 1):
                    cell = s.cell(c.row, col)
                    if cell is not None and cell.is_number and s.is_printed(c.row, col) and c.row not in (s.rows_hidden):
                        if "april 1" in label or "beginning" in label:
                            continue  # opening balances are typed (prior-year signed figures)
                        assert cell.formula is not None, f"{name}!{cell.addr} total is typed"


def test_shared_formulas_and_openpyxl(clean, wb: Workbook) -> None:
    assert sum(1 for s in wb.sheets for c in s.iter_cells() if c.shared_si is not None) > 30
    owb = openpyxl.load_workbook(clean.path)
    assert owb.sheetnames == SHEET_ORDER


def test_evaluator_self_test_is_exact(wb: Workbook) -> None:
    formulas = {(s.name, c.row, c.col): c.formula for s in wb.sheets for c in s.iter_cells() if c.formula}

    def cached(k: tuple[str, int, int]):
        c = wb.cell(*k)
        return None if c is None else c.value

    rc = Recalculator({}, cached, bounds=lambda sh: (wb.sheet(sh).max_row, wb.sheet(sh).max_col))
    mism = []
    for (sheet, r, c), f in formulas.items():
        v = rc.ev.evaluate(f, sheet, r, c)
        if not values_equal(cached((sheet, r, c)), v):
            mism.append((sheet, r, c, f))
    assert not mism, mism[:5]
    assert len(formulas) > 700


def test_typed_narrative_matches_printed_figures(clean, wb: Workbook) -> None:
    s, r, _ = clean.index["n2934.gc.p1"]
    p1 = wb.cell(s, r, 1)
    assert p1.formula is None and isinstance(p1.value, str)
    pat = printed_alias(wb, clean.index, "pl.pat", "CY")
    assert f"{abs(pat):,.2f}" in p1.value
    assert "earned a profit" in p1.value and "previous year: loss" in p1.value
    s, r, _ = clean.index["n2934.gc.p2"]
    p2 = wb.cell(s, r, 1).value
    cash_py = printed_alias(wb, clean.index, "bs.cash", "PY")
    assert f"{cash_py:,.2f}" in p2
    for i in range(4):
        s, r, _ = clean.index[f"n2934.gc.f{i}"]
        assert wb.cell(s, r, 1).value.startswith(f"({'abcd'[i]})")
        assert not wb.sheet(s).row_hidden(r)


def test_ratio_variances_explained(clean, wb: Workbook) -> None:
    for key in ("current_ratio", "debt_equity", "dscr", "roe", "tr_turnover", "tp_turnover", "nct", "npr", "roce"):
        s, r, cols = clean.index[f"note.ratio.{key}"]
        var = wb.cell(s, r, cols["var"]).value
        reason = wb.cell(s, r, cols["reason"]).value
        if abs(var) > 0.25:
            assert reason and reason != "-", key


def test_signature_blocks_and_units(clean, wb: Workbook) -> None:
    blocks = clean.spec.entity["signature_blocks"]
    assert set(blocks) == {"BS", "PL", "SOCIE", "CFS", "Note 29-34"}
    dins = set()
    for sheet, b in blocks.items():
        s = wb.sheet(sheet)
        texts = [s.cell(b["ids_row"], c).value for c in b["cols"][1:]]
        assert all(t.startswith("DIN: ") for t in texts)
        dins.add(tuple(texts))
        assert s.cell(b["date_row"], b["cols"][1]).value == "Date: May 15, 2026"
    assert len(dins) == 1
    for name in ("BS", "PL", "SOCIE", "CFS"):
        assert wb.sheet(name).cell(3, 1).value == "(All amounts in ₹ lakhs, unless otherwise stated)"


def test_templates_recorded(clean) -> None:
    t = clean.spec.templates
    assert t["library"].startswith("house-templates@")
    assert len(t["library_sha256"]) == 64
    assert t["layout.bs"] == "1" and t["wording.gc"] == "1"


def test_exact_mode_natural_rounding(tmp_path) -> None:
    r = cached_workbook(seed=7, exact=True)
    natural = [e for e in r.expected if e.source == "natural"]
    assert natural and all(e.rule_id in ("CAST-01", "BAL-01", "BAL-02") for e in natural)
    assert {(e.sheet, e.cells[0]) for e in natural if e.rule_id == "CAST-01"} == {(i.sheet, i.cell) for i in scan_casting(r.path)}
    for kind in ("bs", "cf"):
        for k in PERIODS:
            assert abs(r.gaps[kind][k]) <= Decimal("0.20")
    assert any(r.gaps[kind][k] != 0 for kind in ("bs", "cf") for k in PERIODS)
    assert simulate_totals_fix(r.path, r.index) == r.gaps


def test_lineage_every_note_line_traces_to_tb(clean) -> None:
    tb_lines = [cs for s in clean.spec.sheets for cs in s.cells.values() if cs.lineage.startswith("tb:")]
    assert len(tb_lines) > 120
    assert all("SUMIF(NS_TB_FT27!" in (cs.formula or "") for cs in tb_lines)
