"""Fix engine: totals-only rewrite, plug removal, OOXML patching, verification, idempotency."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from helpers_core import mini_statements

from fscore import api
from fscore.config import EngagementConfig
from fscore.io.ooxml_patch import PatchError, WorkbookLockedError, XlsxPatcher, validate_package
from fscore.io.xlsx_reader import Workbook
from fscore.models import TextEdit
from fscore.prepare.spec import SheetSpec, WorkbookSpec
from fscore.prepare.writer import write_workbook
from fscore.runner import run_checks

CFG = EngagementConfig()


def _broken(tmp: Path) -> Path:
    def mutate(spec: WorkbookSpec) -> None:
        n = spec.sheet("Note 4")
        n.set(7, 3, formula="SUM(C4:C6)", style="amount_bold")
        n.set(7, 4, formula="SUM(D4:D6)-0.01", style="amount_bold")

    return mini_statements(tmp / "broken.xlsx", mutate)


def test_fix_rewrites_failing_totals_and_verifies(tmp_path: Path) -> None:
    src = _broken(tmp_path)
    before = src.read_bytes()
    res = run_checks(src, CFG)
    assert any(f.rule_id == "CAST-01" for f in res.findings)
    plan = api.propose_fixes(src, CFG, check_result=res)
    edits = {(e.sheet, e.cell): e for e in plan.edits}
    assert edits[("Note 4", "C7")].new_formula == "SUMPRODUCT(ROUND(C4:C6,2))"
    assert edits[("Note 4", "D7")].new_formula == "SUMPRODUCT(ROUND(D4:D6,2))"  # plug dropped
    assert plan.converged and not plan.unresolved
    moved = {(m.sheet, m.cell): m for m in plan.printed_moves}
    assert moved[("Note 4", "C7")].before == "80.67" and moved[("Note 4", "C7")].after == "80.66"
    assert ("BS", "C8") in moved and ("BS", "C10") in moved  # the change flows to the Balance Sheet
    out = tmp_path / "broken - casting corrected.xlsx"
    applied = api.apply_fixes(src, plan, out, config=CFG)
    assert applied.ok, applied.verification.model_dump()
    assert {s.name for s in applied.verification.steps} >= {"xml", "openpyxl", "diff", "selftest", "recheck"}
    assert src.read_bytes() == before  # C1: the original is never modified
    after = run_checks(out, CFG)
    assert [f for f in after.findings if f.rule_id in ("CAST-01", "PLUG-01", "BAL-01")] == []
    wb = openpyxl.load_workbook(out)
    assert wb["Note 4"]["C7"].value == "=SUMPRODUCT(ROUND(C4:C6,2))"
    assert wb["Note 4"].row_dimensions[1].hidden is False


def test_fix_is_idempotent(tmp_path: Path) -> None:
    src = _broken(tmp_path)
    out = tmp_path / "fixed.xlsx"
    plan = api.propose_fixes(src, CFG)
    assert api.apply_fixes(src, plan, out, config=CFG).ok
    again = api.propose_fixes(out, CFG)
    assert [e for e in again.edits if e.fix_kind.value == "auto"] == []


def test_refuses_to_overwrite_source_and_locked_target(tmp_path: Path) -> None:
    src = _broken(tmp_path)
    plan = api.propose_fixes(src, CFG)
    assert not api.apply_fixes(src, plan, src, config=CFG).ok
    out = tmp_path / "locked.xlsx"
    (tmp_path / "~$locked.xlsx").write_bytes(b"lock")
    p = XlsxPatcher(src)
    with pytest.raises(WorkbookLockedError):
        p.write(out)


def test_plan_rejected_when_workbook_changed(tmp_path: Path) -> None:
    src = _broken(tmp_path)
    plan = api.propose_fixes(src, CFG)
    other = mini_statements(tmp_path / "other.xlsx")
    shutil.copy2(other, src)
    res = api.apply_fixes(src, plan, tmp_path / "o.xlsx", config=CFG)
    assert not res.ok and "changed since the plan" in res.message


def _shared_wb(path: Path) -> Path:
    spec = WorkbookSpec()
    s = spec.add(SheetSpec("Note 9", selected=True))
    s.set(1, 1, "9. Other expenses", style="caption_bold")
    s.set(2, 2, "Year ended March 31, 2026", style="header")
    s.set(2, 3, "Year ended March 31, 2025", style="header")
    for i in range(5):
        s.set(3 + i, 1, f"Expense {i}", style="caption")
        s.set(3 + i, 4, 1.0 + i * 0.111, style="amount")
        s.set(3 + i, 2, formula=f"D{3 + i}*1.001", style="amount")  # becomes a shared-formula group
        s.set(3 + i, 3, formula=f"D{3 + i}*0.5", style="amount")
    s.set(8, 1, "Total", style="caption_bold")
    s.set(8, 2, formula="SUM(B3:B7)", style="amount_bold")
    s.set(8, 3, formula="SUM(C3:C7)", style="amount_bold")
    s.print_area = (1, 1, 8, 3)
    return write_workbook(spec, path, shared_formulas=True)


def test_patcher_expands_shared_formula_group(tmp_path: Path) -> None:
    src = _shared_wb(tmp_path / "shared.xlsx")
    with zipfile.ZipFile(src) as z:
        assert b't="shared"' in z.read("xl/worksheets/sheet1.xml")
    p = XlsxPatcher(src)
    p.set_formula("Note 9", "B3", "ROUND(D3*1.001,2)", 1.0)
    out = p.write(tmp_path / "patched.xlsx")
    assert validate_package(out) == []
    with Workbook.open(out) as wb:
        sh = wb.sheet("Note 9")
        assert sh.cell(3, 2).formula == "ROUND(D3*1.001,2)"
        assert [sh.cell(r, 2).formula for r in range(4, 8)] == [f"D{r}*1.001" for r in range(4, 8)]
        assert all(sh.cell(r, 2).shared_si is None for r in range(3, 8))
        assert sh.cell(5, 3).formula == "D5*0.5"


def test_patcher_inserts_new_cell_in_column_order(tmp_path: Path) -> None:
    src = _shared_wb(tmp_path / "shared.xlsx")
    p = XlsxPatcher(src)
    p.set_formula("Note 9", "E8", "SUM(D3:D7)", 6.11)
    p.hide_col("Note 9", 4)
    p.hide_row("Note 9", 2)
    out = p.write(tmp_path / "ins.xlsx")
    assert validate_package(out) == []
    with Workbook.open(out) as wb:
        sh = wb.sheet("Note 9")
        assert sh.cell(8, 5).formula == "SUM(D3:D7)"
        assert 4 in sh.cols_hidden and 2 in sh.rows_hidden
    openpyxl.load_workbook(out)


def test_xml_validation_catches_malformed_formula_tag(tmp_path: Path) -> None:
    """B-T29: the known '</f/>' bug must fail strict validation."""
    src = _shared_wb(tmp_path / "s.xlsx")
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(bad, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</f>", b"</f/>", 1)
            zout.writestr(info, data)
    errs = validate_package(bad)
    assert errs and "sheet1.xml" in errs[0]


def test_text_edits_are_typed_text_only(tmp_path: Path) -> None:
    src = mini_statements(tmp_path / "t.xlsx")
    out = tmp_path / "t2.xlsx"
    refused = api.apply_text_edits(src, [TextEdit(sheet="BS", cell="A2", new_text="=1+1")], out, CFG)
    assert not refused.ok and "C6" in refused.message
    # a cell that holds data used by a total cannot be overwritten with a heading (lesson 15A / B-T22)
    refused2 = api.apply_text_edits(src, [TextEdit(sheet="Note 4", cell="C4", new_text="Heading")], out, CFG)
    assert not refused2.ok and "used by" in refused2.message
    ok = api.apply_text_edits(src, [TextEdit(sheet="BS", cell="A2", new_text="Balance Sheet as at March 31, 2026 ")], out, CFG)
    assert ok.ok, ok.message
    with Workbook.open(out) as wb:
        assert wb.sheet("BS").cell(2, 1).value == "Balance Sheet as at March 31, 2026 "


def test_patch_error_on_missing_sheet(tmp_path: Path) -> None:
    src = _shared_wb(tmp_path / "s.xlsx")
    with pytest.raises(PatchError):
        XlsxPatcher(src).set_formula("Nope", "A1", "1")
